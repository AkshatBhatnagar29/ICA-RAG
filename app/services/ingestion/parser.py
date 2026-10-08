import io
import logging
import subprocess
import time

import fitz  # type: ignore[import-untyped]
from docx import Document as DocxDocument

from app.core.config import Settings
from app.core.exceptions import DocumentParseError, OCRProcessingError
from app.observability.metrics import metrics
from app.services.ingestion.normalization import normalize_text
from app.services.ingestion.processing_models import Block, BlockType, ExtractionMethod, Page
from app.services.ingestion.structure import classify_block, section_for, split_structural_lines

logger = logging.getLogger(__name__)


class OCRProcessor:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def extract(self, image: bytes) -> str:
        started = time.perf_counter()
        metrics.increment("ocr_requests_total")
        if not self.settings.ocr_enabled:
            raise OCRProcessingError("OCR is disabled")
        try:
            result = subprocess.run(
                [self.settings.ocr_executable, "stdin", "stdout", "--psm", "3"],
                input=image,
                capture_output=True,
                timeout=self.settings.ocr_timeout_seconds,
                check=False,
                shell=False,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            metrics.increment("ocr_failures_total", reason="execution")
            raise OCRProcessingError("OCR could not be completed") from exc
        if result.returncode != 0:
            metrics.increment("ocr_failures_total", reason="nonzero_exit")
            raise OCRProcessingError("OCR could not be completed")
        text = normalize_text(result.stdout.decode("utf-8", errors="replace"))
        if not text:
            metrics.increment("ocr_failures_total", reason="empty_output")
            raise OCRProcessingError("OCR returned no text")
        metrics.observe("ocr_latency_ms", (time.perf_counter() - started) * 1000)
        return text


class PDFProcessor:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self.ocr = OCRProcessor(settings)

    def process(self, data: bytes) -> list[Page]:
        try:
            document = fitz.open(stream=data, filetype="pdf")
        except Exception as exc:
            raise DocumentParseError("The PDF could not be parsed") from exc
        pages: list[Page] = []
        try:
            for index, page in enumerate(document):
                pages.append(self._page(page, index + 1))
        finally:
            document.close()
        return pages

    def _page(self, page: fitz.Page, page_number: int) -> Page:
        raw_blocks = page.get_text("blocks", sort=True)
        candidates: list[tuple[str, BlockType, tuple[float, float, float, float] | None]] = []
        for item in raw_blocks:
            text = normalize_text(str(item[4]))
            if not text:
                continue
            bbox = (float(item[0]), float(item[1]), float(item[2]), float(item[3]))
            derived = split_structural_lines(text)
            if len(derived) <= 1:
                candidates.append((text, derived[0][1] if derived else classify_block(text), bbox))
            else:
                candidates.extend((derived_text, derived_type, bbox) for derived_text, derived_type in derived)
        native_text = normalize_text("\n\n".join(item[0] for item in candidates))
        image_only = bool(page.get_images(full=True)) and len(native_text) < self.settings.ocr_min_text_chars
        method = ExtractionMethod.NATIVE_TEXT
        confidence: float | None = None
        if image_only:
            logger.info("ocr_started page_number=%s", page_number)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            native_text = self.ocr.extract(pixmap.tobytes("png"))
            method = ExtractionMethod.OCR
            candidates = [(line, classify_block(line), None) for line in native_text.splitlines() if line.strip()]
            logger.info("ocr_completed page_number=%s", page_number)
        sections = section_for([(text, kind, None) for text, kind, _ in candidates])
        blocks = [
            Block(
                block_id=f"p{page_number}-b{position}",
                # Position alone is not reliable header/footer evidence. Some
                # native PDFs begin their substantive content near the top
                # margin; dropping it would destroy clauses and provenance.
                block_type=kind,
                text=text,
                block_position=position,
                section=sections[position - 1],
                bbox=bbox,
            )
            for position, (text, kind, bbox) in enumerate(candidates, 1)
        ]
        return Page(physical_page_number=page_number, text=normalize_text(native_text),
                    extraction_method=method, blocks=blocks, ocr_confidence=confidence)


class DOCXProcessor:
    def process(self, data: bytes) -> list[Page]:
        try:
            document = DocxDocument(io.BytesIO(data))
        except Exception as exc:
            raise DocumentParseError("The DOCX could not be parsed") from exc
        entries: list[tuple[str, BlockType, str | None]] = []
        for paragraph in document.paragraphs:
            text = normalize_text(paragraph.text)
            if text:
                style_name = paragraph.style.name if paragraph.style is not None else None
                entries.append((text, classify_block(text, style_name), style_name))
        for table_index, table in enumerate(document.tables, 1):
            entries.append((f"Table {table_index}", BlockType.TABLE, None))
            for row_index, row in enumerate(table.rows, 1):
                cells = [normalize_text(cell.text) for cell in row.cells]
                entries.append((" | ".join(cells), BlockType.TABLE_ROW, None))
        sections = section_for(entries)
        blocks = [Block(block_id=f"b{position}", block_type=kind, text=text,
                        block_position=position, section=sections[position - 1])
                  for position, (text, kind, _) in enumerate(entries, 1)]
        text = normalize_text("\n\n".join(block.text for block in blocks))
        # No physical page number is invented for an unrendered DOCX.
        return [Page(physical_page_number=None, text=text, extraction_method=ExtractionMethod.DOCX_NATIVE,
                     blocks=blocks)]
