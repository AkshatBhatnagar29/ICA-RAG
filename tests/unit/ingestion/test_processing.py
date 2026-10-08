import io

from docx import Document

from app.services.ingestion.normalization import normalize_text
from app.services.ingestion.parser import DOCXProcessor
from app.services.ingestion.processing_models import BlockType
from app.services.ingestion.structure import classify_block


def test_normalization_preserves_unicode_and_unwraps_hyphenation() -> None:
    assert normalize_text("termina-\ntion  \n\n\n ₹10,00,000 — 30 days") == "termination\n\n₹10,00,000 — 30 days"
    assert normalize_text("well-known\nterm") == "well-known\nterm"


def test_structure_detection_is_deterministic() -> None:
    assert classify_block("8.2 Limitation of Liability") == BlockType.HEADING
    assert classify_block("- notice period") == BlockType.LIST
    assert classify_block("The parties agree.") == BlockType.PARAGRAPH


def test_docx_parser_retains_heading_clause_unicode_and_table() -> None:
    document = Document()
    document.add_heading("Termination", level=1)
    document.add_paragraph("8.2 Either party may terminate on 30 days' notice.")
    document.add_paragraph("Amount: ₹10,00,000")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Party"
    table.rows[0].cells[1].text = "Notice"
    stream = io.BytesIO()
    document.save(stream)

    page = DOCXProcessor().process(stream.getvalue())[0]

    assert page.physical_page_number is None
    assert any(block.block_type == BlockType.HEADING for block in page.blocks)
    assert any(block.block_type == BlockType.TABLE_ROW for block in page.blocks)
    assert "₹10,00,000" in page.text
    assert page.blocks[1].section == "8.2 Either party may terminate on 30 days' notice."
